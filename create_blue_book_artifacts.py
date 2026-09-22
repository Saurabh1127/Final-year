#!/usr/bin/env python3
"""
create_blue_book_artifacts.py

Generates:
1. SAMVADA_Blue_Book_Report.txt (and Blue_Book_Report.txt) - Clean plain text for copy-pasting
2. SAMVADA_Blue_Book_Report.html - Print-ready HTML formatted with Times New Roman 12pt, 1.5 spacing
3. SAMVADA_Blue_Book_Report.docx (and Blue_Book_Report.docx) - Fully styled Word Document
"""

import os
import re
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def clean_text_for_word(text):
    text = re.sub(r'\*\*(.*?)\*\*', r'\1', text)
    text = re.sub(r'\*(.*?)\*', r'\1', text)
    text = re.sub(r'`(.*?)`', r'\1', text)
    # Remove XML-incompatible control characters
    text = ''.join(c for c in text if ord(c) >= 32 or c in '\n\r\t')
    return text

def set_cell_border(cell, **kwargs):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = tcPr.first_child_found_in("w:tcBorders")
    if tcBorders is None:
        tcBorders = OxmlElement('w:tcBorders')
        tcPr.append(tcBorders)
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        edge_data = kwargs.get(edge)
        if edge_data:
            tag = 'w:{}'.format(edge)
            element = tcBorders.find(qn(tag))
            if element is None:
                element = OxmlElement(tag)
                tcBorders.append(element)
            for key in ["sz", "val", "color", "space", "shadow"]:
                if key in edge_data:
                    element.set(qn('w:{}'.format(key)), str(edge_data[key]))

def generate_txt():
    print("Generating TXT version...")
    with open("SAMVADA_Blue_Book_Report.md", "r", encoding="utf-8") as f:
        content = f.read()

    # Replace explicit page breaks
    content = re.sub(
        r'<div style=\"page-break-before:\s*always;\"></div>',
        '\n\n' + '='*80 + '\n\n',
        content
    )

    # Replace signatures tables
    def table_replacer(match):
        table_html = match.group(0)
        rows = re.findall(r'<tr>(.*?)</tr>', table_html, re.DOTALL)
        out = []
        for r in rows:
            tds = re.findall(r'<td[^>]*>(.*?)</td>', r, re.DOTALL)
            clean_tds = [re.sub(r'<[^>]+>', ' ', td).strip() for td in tds]
            if len(clean_tds) == 2:
                out.append(f"{clean_tds[0]:<45} {clean_tds[1]}")
            else:
                out.append("   ".join(clean_tds))
        return '\n' + '\n'.join(out) + '\n'

    content = re.sub(r'<table.*?>.*?</table>', table_replacer, content, flags=re.DOTALL)

    # Replace Headings
    def h1_repl(m):
        txt = re.sub(r'<[^>]+>', '', m.group(1)).strip()
        return f"\n\n{'='*80}\n{txt.upper()}\n{'='*80}\n"

    def h2_repl(m):
        txt = re.sub(r'<[^>]+>', '', m.group(1)).strip()
        return f"\n\n{'-'*80}\n{txt}\n{'-'*80}\n"

    def h3_repl(m):
        txt = re.sub(r'<[^>]+>', '', m.group(1)).strip()
        return f"\n\n### {txt}\n"

    content = re.sub(r'<h1[^>]*>(.*?)</h1>', h1_repl, content, flags=re.DOTALL)
    content = re.sub(r'<h2[^>]*>(.*?)</h2>', h2_repl, content, flags=re.DOTALL)
    content = re.sub(r'<h3[^>]*>(.*?)</h3>', h3_repl, content, flags=re.DOTALL)

    # Replace breaks and general HTML wrappers
    content = re.sub(r'<br\s*/?>', '\n', content)
    content = re.sub(r'</?(?:p|div|span|b|i|strong|em)(?:\s+[^>]*)?>', '', content)

    # Normalize excessive newlines
    content = re.sub(r'\n{4,}', '\n\n\n', content)

    with open("SAMVADA_Blue_Book_Report.txt", "w", encoding="utf-8") as f:
        f.write(content)
    with open("Blue_Book_Report.txt", "w", encoding="utf-8") as f:
        f.write(content)
    print(f"TXT generated successfully: {len(content.split())} words, {len(content.splitlines())} lines.")

def generate_html():
    print("Generating HTML version...")
    with open("SAMVADA_Blue_Book_Report.md", "r", encoding="utf-8") as f:
        md_text = f.read()

    css = """
    @page {
        size: A4;
        margin: 25mm 20mm 20mm 25mm;
        @bottom-right {
            content: counter(page);
        }
    }
    body {
        font-family: 'Times New Roman', Times, serif;
        font-size: 12pt;
        line-height: 1.5;
        color: #111111;
        background-color: #ffffff;
        margin: 40px auto;
        max-width: 850px;
        padding: 20px;
    }
    h1 {
        font-size: 18pt;
        font-weight: bold;
        text-align: center;
        margin-top: 30px;
        margin-bottom: 20px;
        text-transform: uppercase;
        page-break-before: always;
    }
    h2 {
        font-size: 16pt;
        font-weight: bold;
        margin-top: 24px;
        margin-bottom: 12px;
        color: #000;
        border-bottom: 1px solid #ccc;
        padding-bottom: 4px;
    }
    h3 {
        font-size: 14pt;
        font-weight: bold;
        margin-top: 18px;
        margin-bottom: 8px;
    }
    h4 {
        font-size: 12pt;
        font-weight: bold;
        margin-top: 14px;
        margin-bottom: 6px;
    }
    p {
        text-align: justify;
        margin-bottom: 12px;
        text-indent: 0;
    }
    table {
        width: 100%;
        border-collapse: collapse;
        margin: 18px 0;
        font-size: 11pt;
    }
    table, th, td {
        border: 1px solid #333333;
    }
    th {
        background-color: #f2f2f2;
        font-weight: bold;
        padding: 8px;
        text-align: left;
    }
    td {
        padding: 7px 8px;
        vertical-align: top;
    }
    pre, code {
        font-family: 'Courier New', Courier, monospace;
        background-color: #f8f9fa;
        font-size: 10pt;
    }
    pre {
        padding: 12px;
        border: 1px solid #ddd;
        border-radius: 4px;
        overflow-x: auto;
        white-space: pre-wrap;
        line-height: 1.3;
    }
    blockquote {
        margin: 14px 0;
        padding: 10px 18px;
        background-color: #f8f9fc;
        border-left: 4px solid #3b82f6;
        font-style: italic;
    }
    .image-prompt-box {
        border: 2px dashed #6366f1;
        background-color: #eef2ff;
        padding: 12px 16px;
        margin: 16px 0;
        border-radius: 6px;
        font-size: 10.5pt;
        color: #312e81;
    }
    .diagram-placeholder {
        background-color: #f1f5f9;
        border: 1px solid #cbd5e1;
        padding: 8px 12px;
        font-weight: bold;
        text-align: center;
        margin: 14px 0 4px 0;
    }
    @media print {
        body {
            margin: 0;
            padding: 0;
            max-width: 100%;
        }
        .page-break {
            page-break-before: always;
        }
    }
    """

    html_body = md_text

    def md_table_to_html(match):
        table_text = match.group(0).strip()
        lines = [l.strip() for l in table_text.splitlines() if l.strip()]
        if len(lines) < 2:
            return table_text
        header = lines[0].split('|')[1:-1]
        rows = lines[2:]
        out = ['<table class="report-table">', '  <thead><tr>']
        for h in header:
            out.append(f'    <th>{h.strip()}</th>')
        out.append('  </tr></thead>\n  <tbody>')
        for r in rows:
            cols = r.split('|')[1:-1]
            out.append('  <tr>')
            for c in cols:
                out.append(f'    <td>{c.strip()}</td>')
            out.append('  </tr>')
        out.append('  </tbody>\n</table>')
        return '\n'.join(out)

    table_pattern = re.compile(r'(?:\|[^\n]+\|\n)(?:\|[-: |]+\|\n)(?:\|[^\n]+\|\n?)+')
    html_body = table_pattern.sub(md_table_to_html, html_body)

    html_body = re.sub(
        r'```\n(\[DIAGRAM PLACEHOLDER:.*?\]\nFigure [0-9.]+:[^\n]+)\n```',
        r'<div class="diagram-placeholder">\1</div>',
        html_body
    )

    full_html = f"""<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0">
<title>SAMVADA - Final Year Blue Book Project Report</title>
<style>
{css}
</style>
</head>
<body>
{html_body}
</body>
</html>
"""
    with open("SAMVADA_Blue_Book_Report.html", "w", encoding="utf-8") as f:
        f.write(full_html)
    print("HTML generated successfully.")

def generate_docx():
    print("Generating DOCX version...")
    with open("SAMVADA_Blue_Book_Report.md", "r", encoding="utf-8") as f:
        md_text = f.read()

    doc = docx.Document()

    # Configure Margins (Left 1.25 for binding, Top/Bottom/Right 1.0 inch)
    for s in doc.sections:
        s.top_margin = Inches(1.0)
        s.bottom_margin = Inches(1.0)
        s.left_margin = Inches(1.25)
        s.right_margin = Inches(1.0)
        s.different_first_page_header_footer = True

    # Configure Default Styles
    style_normal = doc.styles['Normal']
    style_normal.font.name = 'Times New Roman'
    style_normal.font.size = Pt(12)
    style_normal.paragraph_format.line_spacing = 1.5
    style_normal.paragraph_format.space_after = Pt(6)

    lines = md_text.splitlines()
    i = 0
    in_code_block = False
    code_lines = []

    while i < len(lines):
        line = lines[i]

        if 'page-break-before: always' in line:
            doc.add_page_break()
            i += 1
            continue

        if line.strip().startswith('```'):
            if not in_code_block:
                in_code_block = True
                code_lines = []
            else:
                in_code_block = False
                code_str = '\n'.join(code_lines)
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(4)
                p.paragraph_format.space_after = Pt(6)
                p.paragraph_format.line_spacing = 1.15
                run = p.add_run(code_str)
                run.font.name = 'Courier New'
                run.font.size = Pt(9.5)
                run.font.color.rgb = RGBColor(40, 40, 40)
            i += 1
            continue

        if in_code_block:
            code_lines.append(line)
            i += 1
            continue

        h1_match = re.search(r'<h1[^>]*>(.*?)</h1>', line, re.IGNORECASE)
        if h1_match:
            doc.add_page_break()
            heading_text = clean_text_for_word(re.sub(r'<[^>]+>', '', h1_match.group(1)).strip())
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(18)
            p.paragraph_format.space_after = Pt(18)
            run = p.add_run(heading_text.upper())
            run.font.name = 'Times New Roman'
            run.font.size = Pt(18)
            run.bold = True
            i += 1
            continue

        h2_match = re.search(r'<h2[^>]*>(.*?)</h2>', line, re.IGNORECASE)
        if h2_match:
            heading_text = clean_text_for_word(re.sub(r'<[^>]+>', '', h2_match.group(1)).strip())
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run(heading_text)
            run.font.name = 'Times New Roman'
            run.font.size = Pt(16)
            run.bold = True
            i += 1
            continue

        h3_match = re.search(r'<h3[^>]*>(.*?)</h3>', line, re.IGNORECASE)
        if h3_match:
            heading_text = clean_text_for_word(re.sub(r'<[^>]+>', '', h3_match.group(1)).strip())
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run(heading_text)
            run.font.name = 'Times New Roman'
            run.font.size = Pt(14)
            run.bold = True
            i += 1
            continue

        if line.startswith('#### '):
            heading_text = clean_text_for_word(line[5:].strip())
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(heading_text)
            run.font.name = 'Times New Roman'
            run.font.size = Pt(12)
            run.bold = True
            i += 1
            continue

        if line.strip().startswith('|') and '|' in line[1:]:
            table_rows = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                row_str = lines[i].strip()
                if not re.match(r'^\|[-: |]+\|$', row_str):
                    cols = [c.strip() for c in row_str.split('|')[1:-1]]
                    table_rows.append(cols)
                i += 1
            
            if table_rows:
                num_cols = max(len(r) for r in table_rows)
                tbl = doc.add_table(rows=len(table_rows), cols=num_cols)
                tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
                for r_idx, row in enumerate(table_rows):
                    for c_idx, cell_value in enumerate(row):
                        if c_idx < num_cols:
                            cell = tbl.cell(r_idx, c_idx)
                            clean_val = clean_text_for_word(cell_value)
                            cell.text = clean_val
                            cp = cell.paragraphs[0]
                            cp.paragraph_format.line_spacing = 1.15
                            cp.paragraph_format.space_after = Pt(2)
                            cp.paragraph_format.space_before = Pt(2)
                            if r_idx == 0:
                                for r in cp.runs:
                                    r.bold = True
                                    r.font.name = 'Times New Roman'
                                    r.font.size = Pt(10.5)
                            else:
                                for r in cp.runs:
                                    r.font.name = 'Times New Roman'
                                    r.font.size = Pt(10)
                            set_cell_border(cell,
                                top=dict(sz=4, val='single', color='444444'),
                                bottom=dict(sz=4, val='single', color='444444'),
                                left=dict(sz=4, val='single', color='444444'),
                                right=dict(sz=4, val='single', color='444444'))
            continue

        if line.strip().startswith('> **IMAGE GENERATION PROMPT'):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            run_lbl = p.add_run(clean_text_for_word(line.strip()[2:]))
            run_lbl.bold = True
            run_lbl.font.color.rgb = RGBColor(60, 20, 140)
            if i + 1 < len(lines) and lines[i+1].strip().startswith('>'):
                prompt_line = clean_text_for_word(lines[i+1].strip()[2:])
                p2 = doc.add_paragraph()
                p2.paragraph_format.space_after = Pt(8)
                run_txt = p2.add_run(prompt_line)
                run_txt.italic = True
                run_txt.font.color.rgb = RGBColor(50, 50, 50)
                i += 1
            i += 1
            continue

        if line.strip().startswith('<table'):
            table_html = []
            while i < len(lines) and '</table>' not in lines[i]:
                table_html.append(lines[i])
                i += 1
            if i < len(lines):
                table_html.append(lines[i])
                i += 1
            full_tbl = '\n'.join(table_html)
            rows = re.findall(r'<tr>(.*?)</tr>', full_tbl, re.DOTALL)
            doc_tbl = doc.add_table(rows=len(rows), cols=2)
            doc_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            for r_idx, r in enumerate(rows):
                tds = re.findall(r'<td[^>]*>(.*?)</td>', r, re.DOTALL)
                for c_idx, td in enumerate(tds[:2]):
                    cell = doc_tbl.cell(r_idx, c_idx)
                    cell.text = re.sub(r'<[^>]+>', '', td).strip()
                    cp = cell.paragraphs[0]
                    cp.paragraph_format.line_spacing = 1.2
                    for r_run in cp.runs:
                        r_run.font.name = 'Times New Roman'
                        r_run.font.size = Pt(11)
            continue

        clean_l = clean_text_for_word(re.sub(r'<[^>]+>', '', line)).strip()
        if clean_l:
            p = doc.add_paragraph()
            p.paragraph_format.line_spacing = 1.5
            p.paragraph_format.space_after = Pt(6)
            if 'text-align: center' in line or clean_l.startswith('[IMAGE PROMPT:') or clean_l.startswith('Project Report') or clean_l.startswith('SAMVADA:') or clean_l == 'on' or clean_l == 'CERTIFICATE' or clean_l == 'PROJECT APPROVAL CERTIFICATE' or clean_l == 'ACKNOWLEDGEMENT' or clean_l.startswith('Blue Book Plagiarism'):
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

            run = p.add_run(clean_l)
            run.font.name = 'Times New Roman'
            run.font.size = Pt(12)
            if clean_l.startswith('SAMVADA:'):
                run.bold = True
                run.font.size = Pt(18)
            elif clean_l in ['CERTIFICATE', 'PROJECT APPROVAL CERTIFICATE', 'ACKNOWLEDGEMENT', 'INDEX', 'LIST OF FIGURES', 'LIST OF TABLES', 'ABSTRACT']:
                run.bold = True
                run.font.size = Pt(16)
            elif clean_l.startswith('Figure ') or clean_l.startswith('Table '):
                run.bold = True
                run.font.size = Pt(11)

        i += 1

    doc.save("SAMVADA_Blue_Book_Report.docx")
    doc.save("Blue_Book_Report.docx")
    print("DOCX generated successfully: SAMVADA_Blue_Book_Report.docx and Blue_Book_Report.docx")

if __name__ == "__main__":
    generate_txt()
    generate_html()
    generate_docx()
